#!/usr/bin/env python3
"""Convert markdown article to Word document."""

from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
import re

def convert_md_to_docx(md_path, docx_path):
    """Convert markdown file to Word document."""

    with open(md_path, 'r', encoding='utf-8') as f:
        content = f.read()

    doc = Document()

    # Set default font
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(12)

    lines = content.split('\n')
    i = 0
    in_table = False
    table_rows = []

    while i < len(lines):
        line = lines[i]

        # Skip empty lines
        if not line.strip():
            i += 1
            continue

        # Handle headers
        if line.startswith('# '):
            p = doc.add_heading(line[2:].strip(), level=1)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif line.startswith('## '):
            doc.add_heading(line[3:].strip(), level=2)
        elif line.startswith('### '):
            doc.add_heading(line[4:].strip(), level=3)
        elif line.startswith('#### '):
            doc.add_heading(line[5:].strip(), level=4)

        # Handle horizontal rules
        elif line.strip() == '---':
            doc.add_paragraph('─' * 50)

        # Handle tables
        elif '|' in line and not line.startswith('$$'):
            # Collect table rows
            table_rows = []
            while i < len(lines) and '|' in lines[i]:
                row = lines[i]
                # Skip separator rows
                if not re.match(r'^\|[-:\s|]+\|$', row):
                    cells = [c.strip() for c in row.split('|')[1:-1]]
                    if cells:
                        table_rows.append(cells)
                i += 1
            i -= 1  # Adjust for the outer loop increment

            if table_rows:
                # Create table
                num_cols = len(table_rows[0])
                table = doc.add_table(rows=len(table_rows), cols=num_cols)
                table.style = 'Table Grid'

                for row_idx, row_data in enumerate(table_rows):
                    for col_idx, cell_text in enumerate(row_data):
                        if col_idx < num_cols:
                            table.rows[row_idx].cells[col_idx].text = cell_text

                doc.add_paragraph()  # Add space after table

        # Handle math blocks (simplified - just add as text)
        elif line.startswith('$$'):
            math_content = line
            i += 1
            while i < len(lines) and not lines[i].strip().endswith('$$'):
                math_content += '\n' + lines[i]
                i += 1
            if i < len(lines):
                math_content += '\n' + lines[i]
            p = doc.add_paragraph(math_content.replace('$$', ''))
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.runs[0].italic = True

        # Handle bold text markers
        elif line.startswith('**') and line.endswith('**'):
            p = doc.add_paragraph()
            run = p.add_run(line.strip('*'))
            run.bold = True

        # Regular paragraphs
        else:
            # Clean up markdown formatting
            text = line
            text = re.sub(r'\*\*(.+?)\*\*', r'\1', text)  # Remove bold markers
            text = re.sub(r'\*(.+?)\*', r'\1', text)  # Remove italic markers
            text = re.sub(r'`(.+?)`', r'\1', text)  # Remove code markers
            doc.add_paragraph(text)

        i += 1

    doc.save(docx_path)
    print(f"Document saved to: {docx_path}")

if __name__ == '__main__':
    convert_md_to_docx(
        '/home/user/ai-research/reports/article_draft_full.md',
        '/home/user/ai-research/reports/article_draft_full.docx'
    )
